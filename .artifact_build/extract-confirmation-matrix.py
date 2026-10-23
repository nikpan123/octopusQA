import hashlib
import json
from pathlib import Path
import openpyxl

source = Path(r'C:\Users\npanek\Downloads\Wymagania biznesowe 26-27.xlsx')
sheet = openpyxl.load_workbook(source, data_only=True).worksheets[0]

def value(row, col):
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            return sheet.cell(merged.min_row, merged.min_col).value
    return sheet.cell(row, col).value

blocks = [(2,8,9,10,'Matematyka','MAT'),(12,18,19,20,'Język polski','MN'),
          (22,28,29,30,'Język polski','MZ'),(32,38,39,40,'Historia','HIS'),
          (42,47,48,49,'Fizyka','FPLUS'),(51,56,57,58,'Fizyka','TNS'),
          (60,65,66,67,'Przyroda','PRZ'),(69,74,75,76,'Biologia','BIO'),
          (78,83,84,85,'Geografia','GEO')]
rows=[]
for start,order,threshold,year,subject,series in blocks:
    for col in range(3,18):
        cls=value(start,col)
        if cls is None or cls=='-': continue
        level=str(value(start+1,col))
        code=value(start+3,col)
        # The empty final class in history/polish is covered by merged cells.
        if code is None or code=='-': raise ValueError((start,col))
        secondary=level.startswith('SŚ')
        scope=('R' if col%2 else 'P') if secondary and series=='MAT' else ''
        rule=str(value(threshold,col))
        order_year='2026/2027' if secondary else str(value(year,col))
        y=int(order_year[:4])
        rows.append(dict(key=f'{series}-{ "SS" if secondary else "SP"}-{cls}-{openpyxl.utils.get_column_letter(col)}',
                         subject=subject,series=series,classLevel=int(cls),level='SŚ' if secondary else 'SP',scope=scope,
                         giftId=int(value(start+2,col)),giftCode=str(code).strip(),
                         orderCode=str(value(order,col)).strip(),coefficient=not secondary and '60%' in rule,
                         orderYear=order_year,start=f'{y}-05-01',end=f'{y+1}-04-30',
                         source=f'{openpyxl.utils.get_column_letter(col)}{start}:{openpyxl.utils.get_column_letter(col)}{year}',
                         block=str(value(start+4,col)),notes=str(sheet.cell(start,1).value)))
out=dict(sourceFile=source.name,sourceSheet=sheet.title,sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
         confirmationYear='2026/2027',clarifications=['SŚ: brak progu ilościowego; wystarczy 1 egzemplarz i pozytywna weryfikacja CC; okres 1.05.2026–30.04.2027. Korekta użytkownika 21.10.2026.',
         'Gratisy SP: aktualny sezon wymaga wartości pola Próg szkoły (współczynnik 60% jest już w nim uwzględniony), bez ponownego mnożenia przez 0,6; dla 400 uczniów minimum wynosi 30. Zachowano warunek więcej niż 5 sztuk. Poprzednie lata: 9 sztuk nie wystarcza, 10 sztuk uruchamia wysyłkę (minimum 10), zgodnie z ostatnim doprecyzowaniem użytkownika. Warunek braku liczby uczniów z A88 pozostaje odrębny.',
         'Wartości scalonych komórek rozwinięto na wszystkie objęte nimi klasy i pozycje.'],rows=rows)
Path('tests/data/potwierdzenia-gratisy-2026.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(f'{len(rows)} pozycji; SP={sum(r["level"]=="SP" for r in rows)}, SŚ={sum(r["level"]=="SŚ" for r in rows)}')
